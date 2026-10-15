# -*- coding: utf-8 -*-
"""把项目里的迁移说明 md 转成排版清晰的 docx（一次性脚本，不入库）。"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1]
DST = sys.argv[2]

EA = "微软雅黑"
MONO = "Consolas"


def set_run_font(run, name=EA, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color


def shade(paragraph, fill="F4F6F8"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def add_rich_text(paragraph, text, size=10.5, base_bold=False):
    """处理 **加粗** 与 `行内代码`"""
    for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        elif token.startswith("`") and token.endswith("`") and len(token) > 1:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=MONO, size=size - 0.5, bold=base_bold,
                         color=RGBColor(0xC0, 0x39, 0x2B))
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=size, bold=base_bold)


doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), EA)

for level, size in ((1, 20), (2, 15), (3, 12.5), (4, 11)):
    st = doc.styles["Heading %d" % level]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x57)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), EA)

with open(SRC, encoding="utf-8") as f:
    lines = f.read().split("\n")

i = 0
n = len(lines)
while i < n:
    line = lines[i]
    stripped = line.strip()

    # 代码块
    if stripped.startswith("```"):
        i += 1
        buf = []
        while i < n and not lines[i].strip().startswith("```"):
            buf.append(lines[i])
            i += 1
        i += 1
        if not buf:
            continue
        # 整块代码放进「一个」段落，行内用 add_break 换行：
        # 上千行代码若每行一个段落，docx 会膨胀到几 MB 且打开很慢。
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        shade(p, "F2F5F7")
        for idx, code_line in enumerate(buf):
            run = p.add_run(code_line if code_line.strip() else " ")
            set_run_font(run, name=MONO, size=8.5)
            if idx < len(buf) - 1:
                run.add_break()
        continue


    # 表格
    if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
        header = [c.strip() for c in stripped.strip("|").split("|")]
        i += 2
        rows = []
        while i < n and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        table = doc.add_table(rows=1, cols=len(header))
        try:
            table.style = "Table Grid"
        except KeyError:
            pass
        hdr = table.rows[0].cells
        for idx, text in enumerate(header):
            hdr[idx].text = ""
            para = hdr[idx].paragraphs[0]
            run = para.add_run(text)
            set_run_font(run, size=10, bold=True)
        for row in rows:
            cells = table.add_row().cells
            for idx in range(len(header)):
                text = row[idx] if idx < len(row) else ""
                cells[idx].text = ""
                para = cells[idx].paragraphs[0]
                add_rich_text(para, text, size=9.5)
        doc.add_paragraph()
        continue

    # 标题
    m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
    if m:
        level = len(m.group(1))
        doc.add_heading(m.group(2).strip(), level=level)
        i += 1
        continue

    # 分隔线
    if stripped == "---":
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        shade(p, "D9DEE3")
        p.text = ""
        i += 1
        continue

    # 引用
    if stripped.startswith(">"):
        text = stripped.lstrip(">").strip()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        add_rich_text(p, text, size=10, base_bold=True)
        i += 1
        continue

    # 空行
    if not stripped:
        i += 1
        continue

    # 列表
    m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
    if m:
        indent = len(m.group(1)) // 2
        text = m.group(3)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18 + indent * 16)
        p.paragraph_format.space_after = Pt(2)
        prefix = "• " if m.group(2) in ("-", "*") else m.group(2) + " "
        run = p.add_run(prefix)
        set_run_font(run, size=10.5, bold=True)
        add_rich_text(p, text, size=10.5)
        i += 1
        continue

    # 普通段落
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    add_rich_text(p, stripped, size=10.5)
    i += 1

doc.save(DST)
print("saved:", DST)
